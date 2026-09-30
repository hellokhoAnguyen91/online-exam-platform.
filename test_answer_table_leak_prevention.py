import io
import docx
from docx_parser import parse_docx_questions, clean_exam_option, extract_answers_from_table, ANSWER_SECTION_HEADER_REGEX

def test_clean_exam_option():
    dirty = 'Giảm lương của người lao động<br><strong>ĐÁP ÁN</strong><br><div class="exam-table-wrapper"><table><tr><td>1</td><td>A</td></tr></table></div>'
    clean = clean_exam_option(dirty)
    assert clean == 'Giảm lương của người lao động', f"Expected clean text, got: {clean}"
    print("✓ test_clean_exam_option passed")

def test_answer_section_header_regex():
    headers = [
        'ĐÁP ÁN', 'ĐÁP ÁN:', 'BẢNG ĐÁP ÁN', 'Bảng tra đáp án', 'HƯỚNG DẪN CHẤM',
        'Key đáp án', 'Answer Key', 'KEY:', 'DANH SÁCH ĐÁP ÁN', 'ĐÁP ÁN CHI TIẾT',
        'PHẦN ĐÁP ÁN'
    ]
    for h in headers:
        assert ANSWER_SECTION_HEADER_REGEX.match(h), f"Regex failed to match: {h}"
    
    # Should NOT match regular question asking for answers
    non_headers = [
        'Đáp án nào sau đây là đúng nhất?',
        'Chọn câu trả lời đúng:',
        'Câu 1: Khái niệm nào sau đây đúng?'
    ]
    for nh in non_headers:
        assert not ANSWER_SECTION_HEADER_REGEX.match(nh), f"Regex should not match: {nh}"
    print("✓ test_answer_section_header_regex passed")

def test_docx_parser_with_trailing_answer_table():
    # Create an in-memory Word document with 2 questions and a trailing answer table
    doc = docx.Document()
    
    # Question 1
    doc.add_paragraph("Câu 1: An toàn lao động là gì?")
    doc.add_paragraph("A. Biện pháp phòng chống tai nạn")
    doc.add_paragraph("B. Quy định về lương")
    doc.add_paragraph("C. Chế độ nghỉ dưỡng")
    doc.add_paragraph("D. Không có đáp án nào")
    
    # Question 2 (The last question)
    doc.add_paragraph("Câu 2: Trách nhiệm của người sử dụng lao động?")
    doc.add_paragraph("A. Huấn luyện an toàn")
    doc.add_paragraph("B. Trang bị bảo hộ")
    doc.add_paragraph("C. Khám sức khỏe định kỳ")
    doc.add_paragraph("D. Cả 3 phương án trên")
    
    # Trailing Answer Section Header
    doc.add_paragraph("ĐÁP ÁN")
    
    # Trailing Answer Table
    table = doc.add_table(rows=3, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Câu"
    hdr_cells[1].text = "Đáp án"
    
    row1_cells = table.rows[1].cells
    row1_cells[0].text = "1"
    row1_cells[1].text = "A"
    
    row2_cells = table.rows[2].cells
    row2_cells[0].text = "2"
    row2_cells[1].text = "D"
    
    # Save to BytesIO
    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    
    # Parse questions
    questions = parse_docx_questions(stream)
    
    assert len(questions) == 2, f"Expected 2 questions, got {len(questions)}"
    
    # Verify Question 1
    q1 = questions[0]
    assert q1["correct_option"] == "A", f"Expected Q1 answer A, got {q1.get('correct_option')}"
    
    # Verify Question 2
    q2 = questions[1]
    assert q2["correct_option"] == "D", f"Expected Q2 answer D, got {q2.get('correct_option')}"
    
    # Ensure Option D of Question 2 does NOT have "ĐÁP ÁN" or table wrapper!
    opt_d = q2["option_d"]
    assert "ĐÁP ÁN" not in opt_d, f"Leak detected in option_d: {opt_d}"
    assert "exam-table-wrapper" not in opt_d, f"Table leak detected in option_d: {opt_d}"
    assert opt_d == "Cả 3 phương án trên", f"Unexpected option_d content: {opt_d}"
    
    print("✓ test_docx_parser_with_trailing_answer_table passed")

if __name__ == "__main__":
    test_clean_exam_option()
    test_answer_section_header_regex()
    test_docx_parser_with_trailing_answer_table()
    print("\nALL ANSWER TABLE LEAK PREVENT TESTS PASSED 100%!")
