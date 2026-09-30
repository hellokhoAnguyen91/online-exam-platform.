import docx
import io
import sys
import docx_parser

if sys.stdout.encoding != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')
if sys.stderr.encoding != 'utf-8':
    sys.stderr.reconfigure(encoding='utf-8')

def run_tests():
    # Test 1: Multi-select with options A-E and (Chọn nhiều đáp án đúng)
    doc = docx.Document()
    doc.add_paragraph('Câu 1. Theo nguyên tắc kiểm soát mối nguy trong an toàn, những biện pháp nào sau đây có thể được sử dụng? (Chọn nhiều đáp án đúng)')
    doc.add_paragraph('A. Triệt tiêu mối nguy')
    doc.add_paragraph('B. Thay thế mối nguy')
    doc.add_paragraph('C. Cô lập mối nguy')
    doc.add_paragraph('D. Chỉ trang bị PPE')
    doc.add_paragraph('E. Thay đổi phương thức quản lý hoạt động')
    doc.add_paragraph('Đáp án: A, B, C, E')

    # Test 2: Essay question
    doc.add_paragraph('Câu 2. Trong quá trình thực hiện LOTO, những nội dung nào sau đây cần được kiểm soát để bảo đảm an toàn? (2 điểm)')

    # Test 3: Multiple choice question WITHOUT any answer indicated
    doc.add_paragraph('Câu 3. Thiết bị bảo hộ cá nhân (PPE) gồm những loại nào sau đây?')
    doc.add_paragraph('A. Mũ bảo hộ')
    doc.add_paragraph('B. Giày bảo hộ')
    doc.add_paragraph('C. Kính bảo hộ')
    doc.add_paragraph('D. Cả A, B và C')

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)

    questions = docx_parser.parse_docx_questions(bio)
    print(f'Total questions: {len(questions)}')

    for i, q in enumerate(questions, 1):
        print(f'--- Q{i} ---')
        print(f"Content: {q['content'][:50]}...")
        print(f"Type: {q['question_type']}")
        print(f"Score: {q.get('score_weight')}")
        print(f"Correct: {repr(q['correct_option'])}")
        print(f"Opt A: {repr(q.get('option_a'))}")
        print(f"Opt D: {repr(q.get('option_d'))}")
        print(f"Opt E: {repr(q.get('option_e'))}")

    assert len(questions) == 3, f"Expected 3 questions, got {len(questions)}"
    
    # Check Q1: Multi-select
    assert questions[0]['question_type'] == 'multi_select', f"Q1 must be multi_select, got {questions[0]['question_type']}"
    assert questions[0]['option_e'] == 'Thay đổi phương thức quản lý hoạt động', f"Q1 option_e wrong: {questions[0]['option_e']}"
    assert 'E.' not in questions[0]['option_d'], f"Q1 option_e must NOT be merged into option_d: {questions[0]['option_d']}"
    assert questions[0]['correct_option'] == 'A,B,C,E', f"Q1 correct_option wrong: {questions[0]['correct_option']}"

    # Check Q2: Essay
    assert questions[1]['question_type'] == 'essay', f"Q2 must be essay, got {questions[1]['question_type']}"
    assert questions[1]['correct_option'] == '', f"Q2 correct_option must be empty, got {repr(questions[1]['correct_option'])}"
    assert questions[1]['option_a'] == '', f"Q2 option_a must be empty"
    assert questions[1]['score_weight'] == 2.0, f"Q2 score_weight wrong: {questions[1]['score_weight']}"

    # Check Q3: Multiple choice WITHOUT answer (must NOT default to 'A'!)
    assert questions[2]['question_type'] == 'multiple_choice', f"Q3 must be multiple_choice"
    assert questions[2]['correct_option'] == '', f"Q3 must have EMPTY correct_option, but got {repr(questions[2]['correct_option'])}"

    print('\n==========================================')
    print('ALL 3 CORE REQUIREMENTS VERIFIED AND PASSED!')
    print('==========================================')

if __name__ == '__main__':
    run_tests()
