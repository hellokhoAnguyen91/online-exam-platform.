import sys
import io
import docx
import docx_parser

sys.stdout.reconfigure(encoding='utf-8')

def test_section_parsing():
    doc = docx.Document()
    doc.add_paragraph("PHẦN I: CÂU HỎI TRẮC NGHIỆM")
    doc.add_paragraph("Câu 1: Thủ đô của Việt Nam là gì?")
    doc.add_paragraph("A. Hà Nội")
    doc.add_paragraph("B. TP. Hồ Chí Minh")
    doc.add_paragraph("C. Đà Nẵng")
    doc.add_paragraph("D. Huế")
    doc.add_paragraph("Đáp án: A")

    doc.add_paragraph("Câu 2: Số nguyên tố chẵn duy nhất là?")
    doc.add_paragraph("A. 0")
    doc.add_paragraph("B. 2")
    doc.add_paragraph("C. 4")
    doc.add_paragraph("D. 6")
    doc.add_paragraph("Đáp án: B")

    doc.add_paragraph("PHẦN II: CÂU HỎI TỰ LUẬN")
    doc.add_paragraph("Câu 3: Hãy trình bày quy trình bảo dưỡng máy móc định kỳ (2.5 điểm).")
    doc.add_paragraph("Câu 4: Nêu các nguyên nhân chính gây mất an toàn điện (1.5 điểm).")

    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)

    qs = docx_parser.parse_docx_questions(bio)
    print("Total parsed:", len(qs))
    for i, q in enumerate(qs, 1):
        print(f"  Q{i}: type={q['question_type']}, score={q['score_weight']}, content={q['content'][:45]}...")

    assert len(qs) == 4, f"Expected 4 questions, got {len(qs)}"
    assert qs[0]["question_type"] == "multiple_choice"
    assert qs[0]["correct_option"] == "A"
    assert qs[1]["question_type"] == "multiple_choice"
    assert qs[1]["correct_option"] == "B"
    assert qs[2]["question_type"] == "essay"
    assert qs[2]["score_weight"] == 2.5
    assert qs[3]["question_type"] == "essay"
    assert qs[3]["score_weight"] == 1.5

    print("\n✓ ALL SECTION PARSER CHECKS PASSED 100%!")

if __name__ == "__main__":
    test_section_parsing()
