import io
import re
import html
from typing import List, Dict, Any, Optional, Tuple
import datetime
import json
import zipfile

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from docx_parser import strip_question_prefix


def clean_html_to_text(html_text: Optional[str]) -> str:
    """Converts HTML text to clean, readable plain text for Excel cells."""
    if not html_text:
        return ""
    text = str(html_text)
    # Replace image tags
    text = re.sub(r'<img[^>]*>', ' [Hình ảnh] ', text, flags=re.IGNORECASE)
    # Replace breaks and paragraphs with newlines
    text = re.sub(r'<(?:br|p|div)[^>]*>', '\n', text, flags=re.IGNORECASE)
    text = re.sub(r'</(?:p|div)>', '\n', text, flags=re.IGNORECASE)
    # Strip any remaining tags
    text = re.sub(r'<[^>]+>', ' ', text)
    # Decode HTML entities (&nbsp;, &amp;, etc.)
    text = html.unescape(text)
    # Normalize whitespace per line
    lines = [re.sub(r'[ \t]+', ' ', line).strip() for line in text.split('\n')]
    # Remove redundant blank lines
    cleaned_lines = []
    for line in lines:
        if line:
            cleaned_lines.append(line)
        elif cleaned_lines and cleaned_lines[-1] != '':
            cleaned_lines.append('')
    return "\n".join(cleaned_lines).strip()


def format_option_text(letter: str, raw_text: Optional[str]) -> str:
    """Formats an option with letter prefix, e.g. 'A. [Nội dung phương án A]'."""
    clean_txt = clean_html_to_text(raw_text)
    if not clean_txt:
        return letter.upper()
    # If the text already starts with 'A.' or 'A:' or 'A -', don't duplicate
    if re.match(r'^[A-F][\s.:)\-–—]', clean_txt, re.IGNORECASE):
        return clean_txt
    return f"{letter.upper()}. {clean_txt}"


def format_answers_text(val: Any, options_map: Dict[str, str], is_essay: bool, is_correct_guide: bool = False, explanation: str = "") -> str:
    """Formats correct answers or student answers for Excel cells."""
    if is_essay:
        if is_correct_guide:
            exp_clean = clean_html_to_text(explanation or val or "")
            if exp_clean:
                return f"Gợi ý đáp án / Rubric chấm:\n{exp_clean}"
            return "Theo thang điểm & hướng dẫn chấm tự luận của giảng viên"
        else:
            txt = clean_html_to_text(val)
            return txt if txt else "(Chưa làm bài tự luận)"

    if val is None or (isinstance(val, str) and not val.strip()):
        return "(Chưa trả lời)" if not is_correct_guide else "(Chưa có đáp án mẫu)"

    # Extract chosen letters
    if isinstance(val, (list, tuple, set)):
        letters = [str(x).strip().upper() for x in val if str(x).strip()]
    elif isinstance(val, str):
        # Extract letters A, B, C, D, E, F
        letters = [c.upper() for c in re.split(r'[,;\s]+', val.strip()) if c]
    else:
        letters = [str(val).strip().upper()]

    if not letters:
        return clean_html_to_text(str(val))

    parts = []
    for let in letters:
        opt_text = options_map.get(let.upper(), "")
        parts.append(format_option_text(let, opt_text))

    return "\n".join(parts) if parts else str(val)


def generate_candidate_excel(
    exam_info: Dict[str, Any],
    candidate_info: Dict[str, Any],
    audit_questions: List[Dict[str, Any]]
) -> bytes:
    """
    Generates an official candidate submission review spreadsheet (.xlsx)
    containing:
      1. STT
      2. Câu hỏi
      3. Câu trả lời đúng
      4. Câu trả lời của thí sinh
      5. Làm đúng (x)
      6. Điểm số (với hàm =SUM ở dòng tổng cộng)
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Chi Tiết Bài Làm"
    ws.views.sheetView[0].showGridLines = True

    # Palette styles
    title_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    alt_row_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    
    font_org = Font(name="Calibri", size=10, bold=True, color="1E3A8A")
    font_main_title = Font(name="Calibri", size=14, bold=True, color="0F172A")
    font_meta = Font(name="Calibri", size=10, color="475569")
    font_candidate = Font(name="Calibri", size=10, bold=True, color="1E293B")
    
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_total_label = Font(name="Calibri", size=11, bold=True, color="0F172A")
    font_total_count = Font(name="Calibri", size=11, bold=True, color="059669")
    font_total_sum = Font(name="Calibri", size=12, bold=True, color="1E3A8A")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )
    total_border = Border(
        top=Side(style="medium", color="1E3A8A"),
        bottom=Side(style="double", color="1E3A8A"),
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1")
    )

    # 1. Title Banner
    ws.merge_cells("A1:F1")
    ws["A1"] = "TRƯỜNG ĐẠI HỌC CÔNG NGHỆ KỸ THUẬT TP. HỒ CHÍ MINH (HCMUTE)"
    ws["A1"].font = font_org
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:F2")
    ws["A2"] = "BẢNG ĐỐI CHIẾU CHI TIẾT BÀI LÀM THÍ SINH"
    ws["A2"].font = font_main_title
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    mc_qs = [q for q in audit_questions if q.get("question_type") != "essay"]
    essay_qs = [q for q in audit_questions if q.get("question_type") == "essay"]

    if not essay_qs:
        eff_mc_max = 10.0
        eff_essay_max = 0.0
    elif not mc_qs:
        eff_mc_max = 0.0
        eff_essay_max = 10.0
    else:
        cfg_mc = float(exam_info.get("mc_max_score") if exam_info.get("mc_max_score") is not None else 7.0)
        cfg_essay = float(exam_info.get("essay_max_score") if exam_info.get("essay_max_score") is not None else 3.0)
        tot_cfg = cfg_mc + cfg_essay
        if tot_cfg <= 0:
            cfg_mc = 7.0
            cfg_essay = 3.0
            tot_cfg = 10.0
        eff_mc_max = round((cfg_mc / tot_cfg) * 10.0, 2)
        eff_essay_max = round(10.0 - eff_mc_max, 2)

    total_mc_w = sum(float(q.get("score_weight") or 1.0) for q in mc_qs)
    mc_scale = (eff_mc_max / total_mc_w) if total_mc_w > 0 else 0.0

    total_essay_w = sum(float(q.get("score_weight") or 1.0) for q in essay_qs)
    essay_scale = (eff_essay_max / total_essay_w) if total_essay_w > 0 else 0.0

    exam_title = exam_info.get("title", "Kỳ thi trắc nghiệm")
    exam_code = exam_info.get("code") or "CK-2026"
    split_info = f" | Thang điểm: Trắc nghiệm {eff_mc_max:g}đ" + (f" + Tự luận {eff_essay_max:g}đ" if essay_qs else "")
    ws.merge_cells("A3:F3")
    ws["A3"] = f"Kỳ thi: {exam_title} | Mã ca thi: {exam_code}{split_info}"
    ws["A3"].font = font_meta
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")

    fullname = candidate_info.get("fullname", "Thí sinh")
    mssv = candidate_info.get("username", "-")
    class_name = candidate_info.get("class_name", "-") or "-"
    submit_str = candidate_info.get("submit_time_str", "-")
    score_display = candidate_info.get("score")
    score_str = f"{score_display}/10.0" if score_display is not None else "Chờ chấm"

    ws.merge_cells("A4:F4")
    ws["A4"] = f"Thí sinh: {fullname} | MSSV: {mssv} | Lớp: {class_name} | Thời gian nộp: {submit_str} | Điểm số: {score_str}"
    ws["A4"].font = font_candidate
    ws["A4"].alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[1].height = 20
    ws.row_dimensions[2].height = 26
    ws.row_dimensions[3].height = 18
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 10

    # 2. Table Headers (Row 6)
    headers = [
        "STT",
        "Nội dung câu hỏi",
        "Câu trả lời đúng",
        "Câu trả lời của thí sinh",
        "Làm đúng (x)",
        "Điểm số"
    ]
    header_row = 6
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.fill = title_fill
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws.row_dimensions[header_row].height = 28

    # 4. Populate Question Rows
    start_data_row = 7
    current_row = start_data_row

    for idx, q in enumerate(audit_questions):
        is_essay = (q.get("question_type") == "essay")
        options_map = {
            "A": q.get("option_a", ""),
            "B": q.get("option_b", ""),
            "C": q.get("option_c", ""),
            "D": q.get("option_d", ""),
            "E": q.get("option_e", ""),
            "F": q.get("option_f", "")
        }

        # Col A: STT
        stt_val = idx + 1
        
        # Col B: Question clean text without 'Câu X.' prefix
        raw_content = q.get("content", "")
        q_cleaned = clean_html_to_text(strip_question_prefix(raw_content))

        # Col C: Correct answer text
        correct_val = q.get("correct") or q.get("correct_option") or ""
        correct_text = format_answers_text(
            val=correct_val,
            options_map=options_map,
            is_essay=is_essay,
            is_correct_guide=True,
            explanation=q.get("explanation", "")
        )

        # Col D: Student answer text
        selected_val = q.get("selected")
        student_text = format_answers_text(
            val=selected_val,
            options_map=options_map,
            is_essay=is_essay,
            is_correct_guide=False
        )

        # Col E: Correct mark 'x'
        # Col F: Numeric points
        earned_pts = 0.0
        mark_val = ""

        if is_essay:
            if q.get("essay_score") is not None:
                essay_sc = float(q.get("essay_score"))
                earned_pts = round(essay_sc * essay_scale, 2)
                mark_val = "x" if essay_sc > 0 else ""
            elif q.get("status") == "pending_grading":
                mark_val = "Chờ chấm"
                earned_pts = 0.0
            else:
                mark_val = ""
                earned_pts = 0.0
        else:
            is_corr = q.get("is_correct") or (q.get("status") == "correct")
            if is_corr:
                mark_val = "x"
                raw_sc = float(q.get("earned_score") if q.get("earned_score") is not None else (q.get("score_weight") or 1.0))
                earned_pts = round(raw_sc * mc_scale, 2)
            elif q.get("status") == "partial":
                mark_val = "x (1 phần)"
                raw_sc = float(q.get("earned_score") or 0.0)
                earned_pts = round(raw_sc * mc_scale, 2)
            else:
                mark_val = ""
                earned_pts = 0.0

        # Write cells
        c_stt = ws.cell(row=current_row, column=1, value=stt_val)
        c_stt.alignment = Alignment(horizontal="center", vertical="top")

        c_q = ws.cell(row=current_row, column=2, value=q_cleaned)
        c_q.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        c_corr = ws.cell(row=current_row, column=3, value=correct_text)
        c_corr.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        c_stu = ws.cell(row=current_row, column=4, value=student_text)
        c_stu.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

        c_mark = ws.cell(row=current_row, column=5, value=mark_val)
        c_mark.alignment = Alignment(horizontal="center", vertical="top")
        if mark_val == "x":
            c_mark.font = Font(name="Calibri", size=12, bold=True, color="059669")
        elif "x" in mark_val:
            c_mark.font = Font(name="Calibri", size=11, bold=True, color="0284C7")
        elif mark_val == "Chờ chấm":
            c_mark.font = Font(name="Calibri", size=10, italic=True, color="D97706")

        c_pts = ws.cell(row=current_row, column=6, value=earned_pts)
        c_pts.alignment = Alignment(horizontal="right", vertical="top")
        c_pts.number_format = "0.00"

        # Borders & Zebra striping
        is_even = ((idx + 1) % 2 == 0)
        for col_idx in range(1, 7):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.border = thin_border
            if is_even and col_idx != 5:
                cell.fill = alt_row_fill

        current_row += 1

    last_data_row = current_row - 1
    if last_data_row < start_data_row:
        last_data_row = start_data_row

    # 5. Bottom Total Row with SUM and COUNTIF
    total_row = current_row
    ws.cell(row=total_row, column=1, value="")
    ws.cell(row=total_row, column=2, value="")
    ws.cell(row=total_row, column=3, value="")
    
    lbl_cell = ws.cell(row=total_row, column=4, value="TỔNG CỘNG:")
    lbl_cell.font = font_total_label
    lbl_cell.alignment = Alignment(horizontal="right", vertical="center")

    # Count of correct questions
    cnt_cell = ws.cell(row=total_row, column=5, value=f'=COUNTIF(E{start_data_row}:E{last_data_row}, "x*")')
    cnt_cell.font = font_total_count
    cnt_cell.alignment = Alignment(horizontal="center", vertical="center")

    # SUM formula for total points
    sum_cell = ws.cell(row=total_row, column=6, value=f"=SUM(F{start_data_row}:F{last_data_row})")
    sum_cell.font = font_total_sum
    sum_cell.alignment = Alignment(horizontal="right", vertical="center")
    sum_cell.number_format = "0.00"

    for col_idx in range(1, 7):
        cell = ws.cell(row=total_row, column=col_idx)
        cell.border = total_border
        cell.fill = total_fill
    ws.row_dimensions[total_row].height = 26

    # 6. Column Widths
    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 46
    ws.column_dimensions["C"].width = 36
    ws.column_dimensions["D"].width = 36
    ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 14

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream.getvalue()


def sanitize_filename(name: str) -> str:
    """Make filename safe for ZIP entry across Windows, Linux, and Mac."""
    if not name:
        return "unnamed"
    v_map = {
        'à':'a','á':'a','ả':'a','ã':'a','ạ':'a','ă':'a','ằ':'a','ắ':'a','ẳ':'a','ẵ':'a','ặ':'a','â':'a','ầ':'a','ấ':'a','ẩ':'a','ẫ':'a','ậ':'a',
        'đ':'d',
        'è':'e','é':'e','ẻ':'e','ẽ':'e','ẹ':'e','ê':'e','ề':'e','ế':'e','ể':'e','ễ':'e','ệ':'e',
        'ì':'i','í':'i','ỉ':'i','ĩ':'i','ị':'i',
        'ò':'o','ó':'o','ỏ':'o','õ':'o','ọ':'o','ô':'o','ồ':'o','ố':'o','ổ':'o','ỗ':'o','ộ':'o','ơ':'o','ờ':'o','ớ':'o','ở':'o','ỡ':'o','ợ':'o',
        'ù':'u','ú':'u','ủ':'u','ũ':'u','ụ':'u','ư':'u','ừ':'u','ứ':'u','ử':'u','ữ':'u','ự':'u',
        'ỳ':'y','ý':'y','ỷ':'y','ỹ':'y','ỵ':'y',
        'À':'A','Á':'A','Ả':'A','Ã':'A','Ạ':'A','Ă':'A','Ằ':'A','Ắ':'A','Ẳ':'A','Ẵ':'A','Ặ':'A','Â':'A','Ầ':'A','Ấ':'A','Ẩ':'A','Ẫ':'A','Ậ':'A',
        'Đ':'D',
        'È':'E','É':'E','Ẻ':'E','Ẽ':'E','Ẹ':'E','Ê':'E','Ề':'E','Ế':'E','Ể':'E','Ễ':'E','Ệ':'E',
        'Ì':'I','Í':'I','Ỉ':'I','Ĩ':'I','Ị':'I',
        'Ò':'O','Ó':'O','Ỏ':'O','Õ':'O','Ọ':'O','Ô':'O','Ồ':'O','Ố':'O','Ổ':'O','Ỗ':'O','Ộ':'O','Ơ':'O','Ờ':'O','Ớ':'O','Ở':'O','Ỡ':'O','Ợ':'O',
        'Ù':'U','Ú':'U','Ủ':'U','Ũ':'U','Ụ':'U','Ư':'U','Ừ':'U','Ứ':'U','Ử':'U','Ữ':'U','Ự':'U',
        'Ỳ':'Y','Ý':'Y','Ỷ':'Y','Ỹ':'Y','Ỵ':'Y'
    }
    res = []
    for ch in str(name):
        res.append(v_map.get(ch, ch))
    clean = "".join(res)
    clean = re.sub(r'[^\w\-_.]', '_', clean)
    clean = re.sub(r'_+', '_', clean).strip('_')
    return clean or "file"


def generate_candidate_audit_excel(res: Any, exam: Any, student: Any, db: Any) -> bytes:
    """Builds full audit data for a single candidate and returns .xlsx bytes."""
    audit_questions = []
    if getattr(res, 'answers_detail', None):
        try:
            audit_questions = json.loads(res.answers_detail)
        except Exception:
            audit_questions = []

    if not audit_questions and getattr(res, 'questions', None):
        try:
            from models import Question
            q_ids = json.loads(res.questions) if isinstance(res.questions, str) else res.questions
            raw_answers = json.loads(res.answers) if (res.answers and isinstance(res.answers, str)) else (res.answers or {})
            questions = db.query(Question).filter(Question.id.in_(q_ids)).all()
            q_dict = {q.id: q for q in questions}
            for q_pos, qid in enumerate(q_ids):
                q = q_dict.get(qid)
                if not q:
                    continue
                selected = raw_answers.get(str(qid)) or raw_answers.get(qid)
                is_corr = (selected == q.correct_option) if selected else False
                st = "correct" if is_corr else ("incorrect" if selected else "unanswered")
                audit_questions.append({
                    "q_idx": q_pos + 1,
                    "id": q.id,
                    "content": q.content,
                    "option_a": q.option_a,
                    "option_b": q.option_b,
                    "option_c": q.option_c,
                    "option_d": q.option_d,
                    "option_e": getattr(q, 'option_e', '') or '',
                    "option_f": getattr(q, 'option_f', '') or '',
                    "question_type": q.question_type or 'multiple_choice',
                    "score_weight": q.score_weight or 1.0,
                    "selected": selected,
                    "correct": q.correct_option,
                    "is_correct": is_corr,
                    "status": st,
                    "earned_score": (q.score_weight or 1.0) if is_corr else 0.0,
                    "explanation": getattr(q, 'explanation', '') or ''
                })
        except Exception as e:
            print(f"Error rebuilding audit for excel: {e}")

    dur_sec = getattr(res, 'duration_seconds', 0) or 0
    dur_str = f"{dur_sec // 60}p {dur_sec % 60:02d}s" if dur_sec > 0 else "-"
    start_str = res.start_time.strftime("%d/%m/%Y %H:%M:%S") if getattr(res, 'start_time', None) else "-"
    submit_str = res.submit_time.strftime("%d/%m/%Y %H:%M:%S") if getattr(res, 'submit_time', None) else "-"
    status_text = "Đã nộp bài" if res.status == "submitted" else "Đang làm dở"
    if res.status == "auto_submitted":
        status_text = "Nộp tự động (Hết giờ)"

    exam_info = {
        "id": exam.id if exam else 1,
        "title": exam.title if exam else "Kỳ thi trắc nghiệm",
        "code": getattr(exam, 'code', '') or '',
        "duration_minutes": getattr(exam, 'duration_minutes', 30),
        "num_questions": getattr(exam, 'num_questions', 10),
        "mc_max_score": getattr(exam, 'mc_max_score', 7.0) if getattr(exam, 'mc_max_score', None) is not None else 7.0,
        "essay_max_score": getattr(exam, 'essay_max_score', 3.0) if getattr(exam, 'essay_max_score', None) is not None else 3.0
    }
    candidate_info = {
        "username": student.username if student else "unknown",
        "fullname": student.fullname if student else "Thí sinh",
        "dob": getattr(student, 'dob', None) or "-",
        "class_name": getattr(student, 'class_name', None) or "-",
        "score": getattr(res, 'score', None),
        "max_score": getattr(res, 'max_score', 10.0) or 10.0,
        "correct_count": getattr(res, 'correct_count', 0) or 0,
        "total_questions": getattr(res, 'total_questions', len(audit_questions)) or len(audit_questions),
        "duration_str": dur_str,
        "start_time_str": start_str,
        "submit_time_str": submit_str,
        "status_text": status_text
    }

    return generate_candidate_excel(exam_info, candidate_info, audit_questions)


def generate_class_summary_excel(exam: Any, results: List[Any], db: Any) -> bytes:
    """Generates the official class-wide summary spreadsheet (.xlsx)."""
    from models import User
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bảng Điểm Tổng Hợp"
    ws.views.sheetView[0].showGridLines = True

    # 1. Header Banners
    font_univ = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
    font_title = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    font_sub = Font(name="Calibri", size=10, italic=True, color="475569")

    ws["A1"] = "TRƯỜNG ĐẠI HỌC CÔNG NGHỆ KỸ THUẬT TP. HỒ CHÍ MINH (HCMUTE)"
    ws["A1"].font = font_univ
    ws["A2"] = "KHOA ĐÀO TẠO & KHẢO THÍ - HỆ THỐNG THI TRỰC TUYẾN"
    ws["A2"].font = font_sub

    exam_title = exam.title if exam else "Kỳ thi trắc nghiệm"
    exam_code = getattr(exam, 'code', '') or ''
    ws["A4"] = f"BẢNG ĐIỂM TỔNG HỢP: {exam_title.upper()}"
    ws["A4"].font = font_title

    now_str = datetime.datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    ws["A5"] = f"Mã kỳ thi: {exam_code} | Thời gian xuất: {now_str} | Tổng số bài nộp: {len(results)}"
    ws["A5"].font = font_sub

    # 2. Table Headers
    headers = [
        "STT", "MSSV", "Họ và tên", "Lớp", "Điểm số (Thang 10)",
        "Số câu đúng", "Tổng số câu", "Tỷ lệ đúng (%)",
        "Thời lượng", "Thời gian bắt đầu", "Thời gian nộp bài", "Trạng thái"
    ]

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    start_row = 7
    ws.row_dimensions[start_row].height = 28
    for col_idx, h_text in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # 3. Data Rows
    current_row = start_row + 1
    for idx, r in enumerate(results):
        user = db.query(User).filter(User.id == r.user_id).first() if db else None
        dur_sec = getattr(r, 'duration_seconds', 0) or 0
        dur_str = f"{dur_sec // 60}p {dur_sec % 60:02d}s" if dur_sec > 0 else "-"
        tot_q = getattr(r, 'total_questions', 0) or 0
        corr_cnt = getattr(r, 'correct_count', 0) or 0
        pct_val = (corr_cnt / tot_q) if (tot_q > 0) else 0.0
        pct_display = f"{round(pct_val * 100, 1)}%" if r.score is not None else "-"

        status_val = "Chờ chấm điểm" if r.score is None else ("Hoàn thành" if r.status == "submitted" else r.status)
        start_t = r.start_time.strftime("%d/%m/%Y %H:%M:%S") if getattr(r, 'start_time', None) else "-"
        submit_t = r.submit_time.strftime("%d/%m/%Y %H:%M:%S") if getattr(r, 'submit_time', None) else "-"

        row_vals = [
            idx + 1,
            user.username if user else f"ID_{r.user_id}",
            user.fullname if user else "Thí sinh",
            getattr(user, 'class_name', '') or "-",
            r.score if r.score is not None else "Chờ chấm",
            corr_cnt if r.score is not None else "Chờ chấm",
            tot_q,
            pct_display,
            dur_str,
            start_t,
            submit_t,
            status_val
        ]

        ws.row_dimensions[current_row].height = 22
        is_even = ((idx + 1) % 2 == 0)
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.border = thin_border
            if is_even:
                cell.fill = alt_fill

            # Alignments
            if col_idx in [1, 2, 4, 6, 7, 8, 9, 10, 11, 12]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 3:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx == 5:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if isinstance(val, (int, float)):
                    cell.number_format = "0.00"
                    if val >= 7.0:
                        cell.font = Font(name="Calibri", size=11, bold=True, color="059669")
                    elif val >= 5.0:
                        cell.font = Font(name="Calibri", size=11, bold=True, color="D97706")
                    else:
                        cell.font = Font(name="Calibri", size=11, bold=True, color="DC2626")

            if col_idx == 2:
                cell.font = Font(name="Calibri", size=11, bold=True, color="1E3A8A")

        current_row += 1

    last_data_row = max(start_row + 1, current_row - 1)

    # 4. Statistics Block
    if len(results) > 0:
        stat_start_row = current_row + 1
        stat_border = Border(
            left=Side(style='thin', color='94A3B8'),
            right=Side(style='thin', color='94A3B8'),
            top=Side(style='thin', color='94A3B8'),
            bottom=Side(style='thin', color='94A3B8')
        )
        stat_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")

        stats = [
            ("Tổng số thí sinh dự thi:", len(results)),
            ("Điểm trung bình cả lớp:", f"=AVERAGE(E{start_row+1}:E{last_data_row})"),
            ("Điểm cao nhất:", f"=MAX(E{start_row+1}:E{last_data_row})"),
            ("Điểm thấp nhất:", f"=MIN(E{start_row+1}:E{last_data_row})"),
            ("Số lượng đạt (>= 5.0 đ):", f'=COUNTIF(E{start_row+1}:E{last_data_row}, ">=5")'),
            ("Tỷ lệ đạt (>= 5.0 đ):", f'=COUNTIF(E{start_row+1}:E{last_data_row}, ">=5")/MAX(1, COUNT(E{start_row+1}:E{last_data_row}))')
        ]

        for s_idx, (s_label, s_val) in enumerate(stats):
            r_idx = stat_start_row + s_idx
            ws.row_dimensions[r_idx].height = 20

            c_lbl = ws.cell(row=r_idx, column=3, value=s_label)
            c_lbl.font = Font(name="Calibri", size=10, bold=True, color="1E3A8A")
            c_lbl.alignment = Alignment(horizontal="right", vertical="center")
            c_lbl.fill = stat_fill
            c_lbl.border = stat_border

            c_val = ws.cell(row=r_idx, column=4, value=s_val)
            c_val.font = Font(name="Calibri", size=10, bold=True, color="0F172A")
            c_val.alignment = Alignment(horizontal="center", vertical="center")
            c_val.fill = stat_fill
            c_val.border = stat_border
            if s_idx in [1, 2, 3]:
                c_val.number_format = "0.00"
            elif s_idx == 5:
                c_val.number_format = "0.0%"

    # 5. Column Widths
    col_widths = {
        "A": 7,
        "B": 15,
        "C": 26,
        "D": 14,
        "E": 20,
        "F": 14,
        "G": 14,
        "H": 16,
        "I": 15,
        "J": 20,
        "K": 20,
        "L": 18
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream.getvalue()


def generate_batch_excel_zip(exam: Any, results: List[Any], db: Any) -> Tuple[bytes, int]:
    """
    Generate an in-memory ZIP archive containing:
      1. '00_Bang_Diem_Tong_Hop_Ca_Lop.xlsx' (Class summary scorecard)
      2. 'Chi_Tiet_Bai_Lam_Tung_Thi_Sinh/BaiThi_[MSSV]_[HoTen].xlsx' for every candidate.
    """
    from models import User
    zip_buffer = io.BytesIO()
    exported_count = 0

    # 1. Summary Excel
    summary_bytes = generate_class_summary_excel(exam, results, db)

    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        # Add summary at root
        zf.writestr("00_Bang_Diem_Tong_Hop_Ca_Lop.xlsx", summary_bytes)

        # 2. Individual candidate detail excels
        for r in results:
            student = db.query(User).filter(User.id == r.user_id).first() if db else None
            username = student.username if student else f"user_{r.user_id}"
            fullname = student.fullname if student else "thi_sinh"

            try:
                cand_bytes = generate_candidate_audit_excel(r, exam, student, db)
                safe_name = sanitize_filename(f"BaiThi_{username}_{fullname}") + ".xlsx"
                zf.writestr(f"Chi_Tiet_Bai_Lam_Tung_Thi_Sinh/{safe_name}", cand_bytes)
                exported_count += 1
            except Exception as e:
                print(f"Error exporting candidate excel for user {username}: {e}")

    zip_buffer.seek(0)
    return zip_buffer.getvalue(), exported_count

